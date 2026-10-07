from collections import defaultdict
from decimal import Decimal
import io
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import generics
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.models import Role, UserProfile
from core.permissions import IsAdminOrOwner, IsCashier

from .models import ShopExpense
from .serializers import ShopExpenseSerializer

CYRILLIC_TO_LATIN = {
    "А": "A", "Б": "B", "В": "V", "Г": "G", "Д": "D", "Е": "E", "Ё": "Yo", "Ж": "J",
    "З": "Z", "И": "I", "Й": "Y", "К": "K", "Л": "L", "М": "M", "Н": "N", "О": "O",
    "П": "P", "Р": "R", "С": "S", "Т": "T", "У": "U", "Ф": "F", "Х": "X", "Ц": "Ts",
    "Ч": "Ch", "Ш": "Sh", "Щ": "Sh", "Ъ": "'", "Ы": "Y", "Ь": "", "Э": "E", "Ю": "Yu", "Я": "Ya",
    "а": "a", "б": "b", "в": "v", "г": "g", "d": "d", "е": "e", "ё": "yo", "ж": "j",
    "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
    "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "x", "ц": "ts",
    "ч": "ch", "ш": "sh", "щ": "sh", "ъ": "'", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
    "ў": "o'", "Ў": "O'", "қ": "q", "Қ": "Q", "ғ": "g'", "Ғ": "G'", "ҳ": "h", "Ҳ": "H",
    "‘": "'", "’": "'", "“": '"', "”": '"', "—": "-", "–": "-", "«": '"', "»": '"',
}

EXPENSE_CAT_NAMES = {
    "RENT": {"uz": "Ijara", "ru": "Аренда", "uz-cyrl": "Ижара"},
    "UTILITIES": {"uz": "Kommunal", "ru": "Коммунальные", "uz-cyrl": "Коммунал"},
    "SUPPLIES": {"uz": "Materiallar / TMTs", "ru": "Материалы / ТМЦ", "uz-cyrl": "Материаллар / ТМЦ"},
    "SALARY": {"uz": "Maosh", "ru": "Зарплата", "uz-cyrl": "Маош"},
    "OTHER": {"uz": "Boshqa", "ru": "Другое", "uz-cyrl": "Бошқа"},
}


def _to_latin_safe(text: str) -> str:
    res = []
    for ch in str(text or ""):
        if ch in CYRILLIC_TO_LATIN:
            res.append(CYRILLIC_TO_LATIN[ch])
        elif ord(ch) < 128:
            res.append(ch)
        else:
            res.append("?")
    return "".join(res)


def _is_admin_or_owner(user) -> bool:
    if getattr(user, "is_superuser", False):
        return True
    profile = getattr(user, "profile", None)
    return isinstance(profile, UserProfile) and profile.role in (Role.ADMIN, Role.OWNER)


class ShopExpensePagination(PageNumberPagination):
    page_size = 25
    page_size_query_param = "page_size"
    max_page_size = 100


class ShopExpenseListCreateView(generics.ListCreateAPIView):
    """
    Owners/admins see all expenses; cashiers see only their own rows.
    Cashiers may POST (withdrawal / outgoing cash).
    """

    serializer_class = ShopExpenseSerializer
    permission_classes = [IsAuthenticated, IsCashier]
    pagination_class = ShopExpensePagination

    def get_queryset(self):
        qs = ShopExpense.objects.select_related("recorded_by").filter(deleted_at__isnull=True)
        if not _is_admin_or_owner(self.request.user):
            qs = qs.filter(recorded_by_id=self.request.user.id)
        from_date = (self.request.query_params.get("from") or "").strip()
        to_date = (self.request.query_params.get("to") or "").strip()
        category = (self.request.query_params.get("category") or "").strip()
        if from_date:
            qs = qs.filter(recorded_at__date__gte=from_date)
        if to_date:
            qs = qs.filter(recorded_at__date__lte=to_date)
        if category:
            qs = qs.filter(category=category)
        return qs.order_by("-recorded_at")


class ShopExpenseDetailView(generics.RetrieveUpdateDestroyAPIView):
    """Owners/admins may edit any expense; cashiers only their own."""

    serializer_class = ShopExpenseSerializer
    permission_classes = [IsAuthenticated, IsCashier]
    lookup_field = "pk"

    def get_queryset(self):
        qs = ShopExpense.objects.select_related("recorded_by").filter(deleted_at__isnull=True)
        if not _is_admin_or_owner(self.request.user):
            qs = qs.filter(recorded_by_id=self.request.user.id)
        return qs

    def destroy(self, request, *args, **kwargs):
        if not _is_admin_or_owner(request.user):
            return Response({"code": "PERMISSION_DENIED"}, status=403)
        instance = self.get_object()
        instance.deleted_at = timezone.now()
        instance.deleted_by = request.user
        instance.save(update_fields=["deleted_at", "deleted_by"])
        return Response(status=204)


class ShiftExpenseSummaryView(APIView):
    """Get expense summary for a specific shift."""
    permission_classes = [IsAuthenticated, IsAdminOrOwner]

    def get(self, request, shift_id):
        from .services import get_shift_summary

        summary = get_shift_summary(shift_id)
        if "error" in summary:
            return Response({"code": "SHIFT_NOT_FOUND", "detail": summary["error"]}, status=404)

        return Response(summary)


class AllShiftsSummaryView(APIView):
    """Get summary of recent shifts with expenses."""
    permission_classes = [IsAuthenticated, IsAdminOrOwner]

    def get(self, request):
        from .services import get_all_shifts_summary

        limit = min(int(request.query_params.get("limit", 10)), 50)
        summaries = get_all_shifts_summary(limit=limit)

        return Response({"results": summaries})


class ShopExpenseExportXlsxView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrOwner]

    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        qs = ShopExpense.objects.select_related("recorded_by").filter(deleted_at__isnull=True)
        from_date = (request.query_params.get("from") or "").strip()
        to_date = (request.query_params.get("to") or "").strip()
        category = (request.query_params.get("category") or "").strip()
        lang = (request.query_params.get("lang") or "uz").strip()

        if from_date:
            qs = qs.filter(recorded_at__date__gte=from_date)
        if to_date:
            qs = qs.filter(recorded_at__date__lte=to_date)
        if category:
            qs = qs.filter(category=category)

        expenses = list(qs.order_by("recorded_at"))

        wb = Workbook()
        ws = wb.active
        ws.title = "Chiqimlar"

        headers = ["№", "Sana va vaqt", "Toifa", "Summa (so'm)", "Kassir / Mas'ul", "Izoh / Sherik"]
        if lang == "ru":
            headers = ["№", "Дата и время", "Категория", "Сумма (сум)", "Кассир / Ответственный", "Примечание / Партнёр"]
        elif lang == "uz-cyrl":
            headers = ["№", "Сана ва вақт", "Тоифа", "Сумма (сўм)", "Кассир / Масъул", "Изоҳ / Шерик"]

        ws.append(headers)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        total_amount = Decimal("0")
        cat_totals = defaultdict(Decimal)

        for idx, exp in enumerate(expenses, start=1):
            cat_label = EXPENSE_CAT_NAMES.get(exp.category, {}).get(lang, exp.category)
            note_str = exp.note or ""

            amt = exp.amount
            total_amount += amt
            cat_totals[cat_label] += amt

            dt_str = timezone.localtime(exp.recorded_at).strftime("%Y-%m-%d %H:%M") if exp.recorded_at else "-"
            ws.append([
                idx,
                dt_str,
                cat_label,
                float(amt),
                exp.recorded_by.username if exp.recorded_by else "-",
                note_str,
            ])

        ws.append([])
        summary_title = "JAMI CHIQIM" if lang == "uz" else ("ИТОГО РАСХОДОВ" if lang == "ru" else "ЖАМИ ЧИҚИМ")
        total_row = ws.max_row + 1
        ws.cell(row=total_row, column=3, value=summary_title).font = Font(bold=True)
        total_cell = ws.cell(row=total_row, column=4, value=float(total_amount))
        total_cell.font = Font(bold=True)
        total_cell.number_format = '#,##0'

        for cat_lbl, cat_amt in cat_totals.items():
            r = ws.max_row + 1
            ws.cell(row=r, column=3, value=f"  · {cat_lbl}")
            c = ws.cell(row=r, column=4, value=float(cat_amt))
            c.number_format = '#,##0'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        date_suffix = f"{from_date}_to_{to_date}" if from_date and to_date else (from_date or to_date or "all")
        filename = f"expenses_{date_suffix}.xlsx"
        resp = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp["Content-Disposition"] = f'attachment; filename="{filename}"'
        return resp


class ShopExpenseExportPdfView(APIView):
    permission_classes = [IsAuthenticated, IsAdminOrOwner]

    def get(self, request):
        qs = ShopExpense.objects.select_related("recorded_by").filter(deleted_at__isnull=True)
        from_date = (request.query_params.get("from") or "").strip()
        to_date = (request.query_params.get("to") or "").strip()
        category = (request.query_params.get("category") or "").strip()
        lang = (request.query_params.get("lang") or "uz").strip()

        if from_date:
            qs = qs.filter(recorded_at__date__gte=from_date)
        if to_date:
            qs = qs.filter(recorded_at__date__lte=to_date)
        if category:
            qs = qs.filter(category=category)

        expenses = list(qs.order_by("recorded_at"))

        lines_txt: list[str] = []
        doc_title = "CHIQIMLAR BO'YICHA HISOBOT" if lang == "uz" else ("OTCHET PO RASKHODAM" if lang == "ru" else "CHIQIMLAR BO'YICHA HISOBOT")
        lines_txt.append(doc_title)
        lines_txt.append("=" * 65)
        period_str = f"Davr: {from_date or 'Boshi'} - {to_date or 'Bugun'}"
        lines_txt.append(_to_latin_safe(period_str))
        lines_txt.append(f"Chop etildi: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}")
        lines_txt.append("-" * 65)

        total_amount = Decimal("0")
        cat_totals = defaultdict(Decimal)

        lines_txt.append(f"{'No':<4} {'Sana':<17} {'Toifa':<18} {'Summa (so\'m)':>14} {'Kassir':<10}")
        lines_txt.append("-" * 65)

        for idx, exp in enumerate(expenses, start=1):
            cat_label = EXPENSE_CAT_NAMES.get(exp.category, {}).get("uz", exp.category)
            amt = exp.amount
            total_amount += amt
            cat_totals[cat_label] += amt
            dt_str = timezone.localtime(exp.recorded_at).strftime("%Y-%m-%d %H:%M") if exp.recorded_at else "-"
            cashier_name = exp.recorded_by.username if exp.recorded_by else "-"
            amt_fmt = f"{amt:,.0f}".replace(",", " ")
            lines_txt.append(f"{idx:<4} {dt_str:<17} {_to_latin_safe(cat_label)[:17]:<18} {amt_fmt:>14} {_to_latin_safe(cashier_name)[:9]:<10}")
            if exp.note:
                note_clean = _to_latin_safe(exp.note)
                lines_txt.append(f"     Izoh: {note_clean[:55]}")

        lines_txt.append("=" * 65)
        total_fmt = f"{total_amount:,.0f}".replace(",", " ")
        lines_txt.append(f"JAMI CHIQIM: {total_fmt:>49} so'm")
        lines_txt.append("-" * 65)
        lines_txt.append("Toifalar kesimida:")
        for cat_lbl, cat_amt in cat_totals.items():
            cat_fmt = f"{cat_amt:,.0f}".replace(",", " ")
            lines_txt.append(f"  - {_to_latin_safe(cat_lbl):<30} {cat_fmt:>20} so'm")

        lines_txt.append("")
        lines_txt.append("Mas'ul shaxs (imzo): ____________________")

        pdf_lines = []
        y = 800
        for row in lines_txt:
            safe = row.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            pdf_lines.append(f"BT /F1 9 Tf 40 {y} Td ({safe[:95]}) Tj ET")
            y -= 13
            if y < 40:
                break
        content_stream = "\n".join(pdf_lines)
        objects = []
        objects.append("1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n")
        objects.append("2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n")
        objects.append(
            "3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>endobj\n"
        )
        objects.append(
            f"4 0 obj<< /Length {len(content_stream)} >>stream\n{content_stream}\nendstream\nendobj\n"
        )
        objects.append("5 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>endobj\n")
        out = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for obj in objects:
            offsets.append(len(out))
            out.extend(obj.encode("latin-1", errors="replace"))
        xref_pos = len(out)
        out.extend(f"xref\n0 {len(offsets)}\n".encode())
        out.extend(b"0000000000 65535 f \n")
        for off in offsets[1:]:
            out.extend(f"{off:010d} 00000 n \n".encode())
        out.extend(
            f"trailer<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
        )
        date_suffix = f"{from_date}_to_{to_date}" if from_date and to_date else (from_date or to_date or "all")
        resp = HttpResponse(bytes(out), content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="expenses_{date_suffix}.pdf"'
        return resp
